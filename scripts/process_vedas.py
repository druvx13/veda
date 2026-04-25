#!/usr/bin/env python3
from __future__ import annotations

import sqlite3
from pathlib import Path
from typing import Iterable, List, Sequence

from openpyxl import Workbook, load_workbook


ROOT = Path(__file__).resolve().parents[1]
SOURCE_XLSX = ROOT / "FourVedas20200922.xlsx"
RAW_EXCEL_DIR = ROOT / "raw" / "excel"
INDIVIDUAL_SQL_DIR = ROOT / "raw" / "sqls" / "individual"
COMBINED_SQL_DIR = ROOT / "raaw" / "sqls"

VEDA_SHEETS: Sequence[str] = ("Rik", "Yaju", "Saam", "Atharva")


def sql_ident(name: str) -> str:
    return '"' + name.replace('"', '""') + '"'


def normalize_headers(headers: Sequence[object]) -> List[str]:
    seen: dict[str, int] = {}
    normalized: List[str] = []
    for index, value in enumerate(headers, start=1):
        text = str(value).strip() if value is not None else ""
        if not text:
            text = f"column_{index}"
        if text in seen:
            seen[text] += 1
            text = f"{text}_{seen[text]}"
        else:
            seen[text] = 1
        normalized.append(text)
    return normalized


def worksheet_to_sqlite(conn: sqlite3.Connection, sheet_name: str, rows: Iterable[Sequence[object]]) -> None:
    iterator = iter(rows)
    try:
        first_row = next(iterator)
    except StopIteration:
        return

    headers = normalize_headers(list(first_row))
    conn.execute(f"DROP TABLE IF EXISTS {sql_ident(sheet_name)}")
    create_columns = ", ".join(f"{sql_ident(col)} TEXT" for col in headers)
    conn.execute(f"CREATE TABLE {sql_ident(sheet_name)} ({create_columns})")

    placeholders = ", ".join("?" for _ in headers)
    insert_sql = f"INSERT INTO {sql_ident(sheet_name)} VALUES ({placeholders})"

    batch: List[Sequence[object]] = []
    for row in iterator:
        values = list(row)
        if len(values) < len(headers):
            values.extend([None] * (len(headers) - len(values)))
        elif len(values) > len(headers):
            values = values[: len(headers)]
        batch.append([None if v is None else str(v) for v in values])
        if len(batch) >= 2000:
            conn.executemany(insert_sql, batch)
            batch.clear()
    if batch:
        conn.executemany(insert_sql, batch)

    conn.commit()


def workbook_to_sqlite(xlsx_path: Path, sqlite_path: Path, include_sheets: Sequence[str] | None = None) -> None:
    sqlite_path.parent.mkdir(parents=True, exist_ok=True)
    if sqlite_path.exists():
        sqlite_path.unlink()

    wb = load_workbook(filename=xlsx_path, read_only=True, data_only=True)
    selected_sheets = include_sheets if include_sheets is not None else wb.sheetnames
    with sqlite3.connect(sqlite_path) as conn:
        for sheet_name in selected_sheets:
            ws = wb[sheet_name]
            worksheet_to_sqlite(conn, sheet_name, ws.iter_rows(values_only=True))
    wb.close()


def split_into_individual_excels() -> list[Path]:
    RAW_EXCEL_DIR.mkdir(parents=True, exist_ok=True)
    wb = load_workbook(filename=SOURCE_XLSX, read_only=True, data_only=True)
    output_paths: list[Path] = []

    for sheet_name in VEDA_SHEETS:
        ws = wb[sheet_name]
        output_path = RAW_EXCEL_DIR / f"{sheet_name.lower()}_veda.xlsx"
        out_wb = Workbook(write_only=True)
        out_ws = out_wb.create_sheet(title=sheet_name)
        for row in ws.iter_rows(values_only=True):
            out_ws.append(list(row))
        out_wb.save(output_path)
        out_wb.close()
        output_paths.append(output_path)

    wb.close()
    return output_paths


def main() -> None:
    individual_files = split_into_individual_excels()

    INDIVIDUAL_SQL_DIR.mkdir(parents=True, exist_ok=True)
    for xlsx_path in individual_files:
        sqlite_name = f"{xlsx_path.stem}.sqlite"
        workbook_to_sqlite(xlsx_path, INDIVIDUAL_SQL_DIR / sqlite_name)

    COMBINED_SQL_DIR.mkdir(parents=True, exist_ok=True)
    workbook_to_sqlite(SOURCE_XLSX, COMBINED_SQL_DIR / "four_vedas.sqlite")


if __name__ == "__main__":
    main()
